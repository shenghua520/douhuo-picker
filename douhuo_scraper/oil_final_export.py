"""油类最终交付表：
- 补完运费 + 拉 sku_list（缓存续跑）
- 多 SKU 拆行
- 组合/礼盒/带米/调味品 → 单独分表，其它表不含
- 删指定字段；采集时间写入文件名
- 排序：规格(体积) → 代发价
"""

from __future__ import annotations

import json
import re
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"D:\Users\shenghua\Downloads\商品筛选\douhuo_scraper")
from client import DouhuoClient  # noqa: E402
from excelio import FIELD_ZH, flatten_value, union_keys  # noqa: E402

PROFILE = Path(r"D:\Users\shenghua\Downloads\商品筛选\state\browser_profile")
SRC = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\油类\油类全量商品_京东类目版.xlsx")
OUT_DIR = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\油类")
CACHE = OUT_DIR / "detail_cache.json"
STATUS = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\progress_status.txt")

HEADER_FILL = PatternFill("solid", fgColor="1D1F27")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LABEL_FILL = PatternFill("solid", fgColor="3D4450")
LABEL_FONT = Font(color="FFFFFF", size=10)

DROP_FIELDS = {
    "goods_id", "default_sku", "supply_type", "status", "main_img",
    "匹配关键词", "采集时间", "sort", "is_update", "update_time",
    "is_select", "is_test",
}

PREFERRED = [
    "子类", "子类细分", "品牌", "规格",
    "spu_sn", "spu_name", "SKU编码", "SKU规格说明",
    "supply_name", "plat_price", "retail_price", "profit",
    "包邮结论", "is_freight", "不包邮_不发货地区", "包邮地区",
    "min_reduce", "max_reduce", "商品链接",
]

ZH_EXTRA = {
    "子类": "标准子类(京东/淘宝)",
    "子类细分": "细分(浓香/低芥酸等)",
    "品牌": "品牌",
    "规格": "规格(从名称解析)",
    "spu_sn": "商品编码",
    "spu_name": "商品名称",
    "SKU编码": "SKU编码 sku_sn",
    "SKU规格说明": "SKU规格属性",
    "supply_name": "渠道名称",
    "plat_price": "代发价",
    "retail_price": "零售价",
    "profit": "利润率(%)",
    "包邮结论": "包邮说明",
    "is_freight": "原始包邮值 1=包邮 2=不包邮/按地区",
    "不包邮_不发货地区": "不包邮/不发货地区",
    "包邮地区": "包邮地区",
    "min_reduce": "集采价下限",
    "max_reduce": "集采价上限",
    "商品链接": "商品链接",
}

COMBO_PAT = re.compile(
    r"(组合|套装|礼盒|套组|套餐|大礼包|大礼盒|"
    r"大米|香米|丝苗|五常|面粉|挂面|面条|米粉|麦片|燕麦|"
    r"调味|酱油|生抽|老抽|蚝油|醋|料酒|鸡精|味精|酱|"
    r"花生油.{0,6}\+|玉米油.{0,6}\+|\+.{0,8}油|"
    r"油.{0,4}\+.{0,8}(米|面|酱|醋)|"
    r"\d\+\d|\d\+\d\+\d)"
)


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def write_status(msg: str) -> None:
    try:
        STATUS.parent.mkdir(parents=True, exist_ok=True)
        STATUS.write_text(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n", encoding="utf-8")
    except OSError:
        pass


def is_combo(name: str) -> bool:
    n = name or ""
    return bool(COMBO_PAT.search(n))


def spec_sort_key(spec: str):
    """按体积/重量排序：先 L/ml 换算毫升，再 kg/g，无法识别放最后。"""
    s = str(spec or "")
    m = re.match(r"(\d+(?:\.\d+)?)\s*(ml|L|kg|g|斤)", s, re.I)
    if not m:
        return (2, 0)
    num = float(m.group(1))
    unit = m.group(2).lower()
    if unit == "l":
        return (0, num * 1000)
    if unit == "ml":
        return (0, num)
    if unit == "kg":
        return (1, num * 1000)
    if unit == "g":
        return (1, num)
    return (2, num)


def norm_regions(raw) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s or s.lower() in {"null", "none", "无"}:
        return ""
    s = s.strip("-")
    parts = [p for p in re.split(r"[-,，、;；/|]+", s) if p and p != "无"]
    return "、".join(parts)


def summarize(is_freight, no_deliver: str, free_ship: str) -> str:
    nd = (no_deliver or "").strip()
    fs = (free_ship or "").strip()
    try:
        freight = int(is_freight) if is_freight not in (None, "") else None
    except (TypeError, ValueError):
        freight = None
    if freight == 1 and not nd:
        return "全国包邮（不包邮地区无）"
    if freight == 1 and nd:
        return f"包邮，但以下地区除外：{nd}"
    if freight == 2 and fs and not nd:
        return f"仅以下地区包邮：{fs}；其余地区不包邮"
    if freight == 2 and nd and not fs:
        return f"不包邮/限运地区：{nd}"
    if freight == 2 and nd and fs:
        return f"包邮地区：{fs}；不包邮/不发货：{nd}"
    if freight == 2:
        return "不包邮（接口未返回地区明细，可能按收货地址计运费）"
    if nd:
        return f"限运地区：{nd}"
    if fs:
        return f"包邮地区：{fs}"
    return "未知（无运费地区信息）"


def sku_attr_text(attr) -> str:
    if not attr:
        return ""
    if isinstance(attr, list):
        parts = []
        for a in attr:
            if isinstance(a, dict):
                parts.append(f"{a.get('name') or ''}:{a.get('val') or ''}".strip(":"))
            else:
                parts.append(str(a))
        return "；".join(p for p in parts if p and p != ":")
    return str(attr)


def dump(ws, rows, headers):
    for col, key in enumerate(headers, start=1):
        c1 = ws.cell(1, col, key)
        c1.fill = HEADER_FILL
        c1.font = HEADER_FONT
        c2 = ws.cell(2, col, ZH_EXTRA.get(key) or FIELD_ZH.get(key) or key)
        c2.fill = LABEL_FILL
        c2.font = LABEL_FONT
        ws.column_dimensions[get_column_letter(col)].width = min(max(len(key) + 2, 12), 28)
    for r_i, row in enumerate(rows, start=3):
        for c_i, key in enumerate(headers, start=1):
            ws.cell(r_i, c_i, flatten_value(row.get(key)))
    ws.freeze_panes = "A3"


def main() -> int:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    wb = load_workbook(SRC, data_only=True)
    ws = wb["食用油全量"]
    raw = list(ws.iter_rows(values_only=True))
    hin = list(raw[0])
    idx = {h: i for i, h in enumerate(hin)}
    base = []
    for r in raw[2:]:
        if not r or r[idx["spu_sn"]] in (None, "") and r[idx["spu_name"]] in (None, ""):
            # 仍用 spu 名称
            pass
        if not r or all(c is None or str(c).strip() == "" for c in r[:6]):
            continue
        row = {h: (r[idx[h]] if h in idx and idx[h] < len(r) else None) for h in hin}
        if not row.get("spu_name"):
            continue
        base.append(row)
    wb.close()
    log(f"读入 SPU {len(base)}")

    cache: dict[str, dict] = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
            log(f"详情缓存 {len(cache)}")
        except json.JSONDecodeError:
            cache = {}

    # 兼容旧 freight_cache
    old_fc = OUT_DIR / "freight_cache.json"
    if old_fc.exists():
        try:
            old = json.loads(old_fc.read_text(encoding="utf-8"))
            for gid, info in old.items():
                if gid not in cache:
                    cache[gid] = {
                        "is_freight": info.get("is_freight"),
                        "no_deliver": info.get("no_deliver"),
                        "free_shipping": info.get("free_shipping"),
                        "sku_list": None,
                    }
            log(f"合并旧运费缓存后 {len(cache)}")
        except json.JSONDecodeError:
            pass

    t0 = time.time()
    last = t0
    with DouhuoClient(PROFILE, headless=True) as client:
        log("登录态 " + client.login_state())
        client._page.wait_for_timeout(3000)
        for i, row in enumerate(base, 1):
            gid = str(row.get("goods_id") or "")
            # goods_id 可能已被删列，从链接取
            if not gid:
                m = re.search(r"id=(\d+)", str(row.get("商品链接") or ""))
                gid = m.group(1) if m else ""
            need_detail = (gid not in cache) or (cache.get(gid, {}).get("sku_list") is None)
            if need_detail and gid:
                d = {}
                try:
                    d = client.post("/manager/Choice/getGoodsDetail", {"goods_id": gid}) or {}
                except Exception as exc:  # noqa: BLE001
                    log(f"  detail fail {gid}: {exc}")
                free_ship = ""
                try:
                    free_ship = client.post("/manager/Choice/free_shipping", {"goods_id": gid})
                except Exception as exc:  # noqa: BLE001
                    pass
                prev = cache.get(gid, {})
                cache[gid] = {
                    "is_freight": (d.get("is_freight") if isinstance(d, dict) else None) or prev.get("is_freight") or row.get("is_freight"),
                    "no_deliver": norm_regions(d.get("no_deliver") if isinstance(d, dict) else prev.get("no_deliver")),
                    "free_shipping": norm_regions(free_ship or prev.get("free_shipping")),
                    "sku_list": d.get("sku_list") if isinstance(d, dict) else prev.get("sku_list"),
                    "goods_sn": (d.get("goods_sign") or d.get("goods_sn")) if isinstance(d, dict) else prev.get("goods_sn"),
                }
                if i % 40 == 0:
                    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
            now = time.time()
            if now - last >= 300 or i == len(base):
                eta = (len(base) - i) / (i / max(now - t0, 1e-6)) / 60 if i else 0
                msg = f"[进度] 详情/SKU {i}/{len(base)} ({i*100//len(base)}%) ETA≈{eta:.0f}分钟"
                log(msg)
                write_status(msg)
                last = now
            elif i % 100 == 0:
                log(f"  … {i}/{len(base)}")
        CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    # 展开 SKU + 组合分类
    out_rows = []
    for row in base:
        gid = str(row.get("goods_id") or "")
        if not gid:
            m = re.search(r"id=(\d+)", str(row.get("商品链接") or ""))
            gid = m.group(1) if m else ""
        info = cache.get(gid, {})
        freight = info.get("is_freight", row.get("is_freight"))
        nd = info.get("no_deliver") or ""
        fs = info.get("free_shipping") or ""
        name = str(row.get("spu_name") or "")
        combo = is_combo(name)
        common = {
            "子类": row.get("子类"),
            "子类细分": row.get("子类细分"),
            "品牌": row.get("品牌"),
            "规格": row.get("规格"),
            "spu_sn": info.get("goods_sn") or row.get("spu_sn"),
            "spu_name": name,
            "supply_name": row.get("supply_name"),
            "plat_price": row.get("plat_price"),
            "retail_price": row.get("retail_price"),
            "profit": row.get("profit"),
            "包邮结论": summarize(freight, nd, fs),
            "is_freight": freight,
            "不包邮_不发货地区": nd,
            "包邮地区": fs,
            "min_reduce": row.get("min_reduce"),
            "max_reduce": row.get("max_reduce"),
            "商品链接": row.get("商品链接"),
            "_combo": combo,
        }
        skus = info.get("sku_list") or []
        if not skus:
            common["SKU编码"] = ""
            common["SKU规格说明"] = ""
            out_rows.append(common)
            continue
        for sk in skus:
            if not isinstance(sk, dict):
                continue
            r2 = dict(common)
            r2["SKU编码"] = sk.get("sku_sn") or sk.get("sku_sign") or ""
            r2["SKU规格说明"] = sku_attr_text(sk.get("attr"))
            if sk.get("plat_price") not in (None, ""):
                r2["plat_price"] = sk.get("plat_price")
            if sk.get("retail_price") not in (None, ""):
                r2["retail_price"] = sk.get("retail_price")
            if sk.get("profit") not in (None, ""):
                r2["profit"] = sk.get("profit")
            if sk.get("sku_name"):
                # 保留 SPU 名，SKU 名放说明补充
                if r2["SKU规格说明"]:
                    r2["SKU规格说明"] = f"{r2['SKU规格说明']} | {sk.get('sku_name')}"
                else:
                    r2["SKU规格说明"] = str(sk.get("sku_name"))
            out_rows.append(r2)

    headers = [h for h in PREFERRED]
    # 排序：规格 → 代发价
    def sk(r):
        return (0 if not r.get("_combo") else 1, spec_sort_key(r.get("规格")), -float(r.get("plat_price") or 0), str(r.get("spu_name") or ""))

    out_rows.sort(key=sk)
    singles = [r for r in out_rows if not r.get("_combo")]
    combos = [r for r in out_rows if r.get("_combo")]

    out_path = OUT_DIR / f"油类商品_京东类目_采集时间_{stamp}.xlsx"
    wb_out = Workbook()
    ws0 = wb_out.active
    ws0.title = "字段说明"
    ws0.append(["字段", "说明"])
    for c in ws0[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for a, b in [
        ("文件", f"采集时间 {stamp}（已写入文件名）"),
        ("组合分表", "「组合_礼盒_套装」：名称含组合/套装/礼盒，或带米/调味品/多油种拼搭等"),
        ("其它分表", "仅单品油，已移除全部组合商品"),
        ("排序", "规格(体积 ml/L)升序 → 代发价降序"),
        ("SKU", "同一商品多 SKU 已拆成多行，见 SKU编码 / SKU规格说明"),
        ("包邮结论", "1 且无地区限制=全国包邮；2 且仅列地区=仅该地区包邮；2 无明细=不包邮按地址计"),
        ("已删除字段", "goods_id/default_sku/渠道类型/上下架/主图/匹配关键词/采集时间/推荐指数/是否更新中/更新时间/是否已选/是否测试"),
        ("is_freight", "1=包邮 2=不包邮或按地区"),
        ("plat_price", "代发价（元）"),
        ("retail_price", "零售价（元）"),
        ("profit", "利润率(%)"),
    ]:
        ws0.append([a, b])
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 78

    def clean(rows):
        return [{k: r.get(k) for k in headers} for r in rows]

    dump(wb_out.create_sheet("食用油单品全量"), clean(singles), headers)
    dump(wb_out.create_sheet("组合_礼盒_套装"), clean(combos), headers)

    by = {}
    for r in singles:
        by.setdefault(r.get("子类") or "未分类", []).append(r)
    JD_ORDER = [
        "菜籽油", "花生油", "玉米油", "大豆油", "葵花籽油", "橄榄油", "调和油",
        "山茶油", "芝麻油", "亚麻籽油", "稻米油", "核桃油", "牛油果油",
        "椰子油", "棕榈油", "红花籽油", "葡萄籽油", "小麦胚芽油",
        "紫苏籽油", "南瓜籽油", "猪油", "功能性食用油", "食用油组合/未细分",
    ]
    used = set()
    for sub in JD_ORDER:
        if sub in by:
            dump(wb_out.create_sheet(re.sub(r'[\\/*?:\[\]]', "_", sub)[:31]), clean(by[sub]), headers)
            used.add(sub)
    for sub, items in by.items():
        if sub not in used:
            dump(wb_out.create_sheet(re.sub(r'[\\/*?:\[\]]', "_", sub)[:31]), clean(items), headers)

    # 统计
    ws_s = wb_out.create_sheet("统计")
    ws_s.append(["项", "数量"])
    for c in ws_s[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    ws_s.append(["单品行数(含SKU拆分)", len(singles)])
    ws_s.append(["组合/礼盒行数(含SKU拆分)", len(combos)])
    multi = Counter()
    for r in out_rows:
        pass
    # 原SPU数
    ws_s.append(["原SPU数", len(base)])
    ws_s.append(["采集时间", stamp])
    for sub, items in sorted(by.items(), key=lambda x: -len(x[1])):
        ws_s.append([f"单品-{sub}", len(items)])
    ws_s.column_dimensions["A"].width = 28
    ws_s.column_dimensions["B"].width = 18

    wb_out.save(out_path)
    log("=" * 40)
    log(f"单品 {len(singles)} 行  组合 {len(combos)} 行  → {out_path.name}")
    log("组合判定关键词: 组合/套装/礼盒/大米/米/调味/酱油/醋/面粉/挂面/+多油种/2+1 等")
    log("子类(单品): " + ", ".join(f"{k}:{len(v)}" for k, v in list(sorted(by.items(), key=lambda x: -len(x[1])))[:12]))
    write_status(f"完成 {out_path.name} 单品{len(singles)} 组合{len(combos)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
