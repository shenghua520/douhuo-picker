"""为油类全量表补充运费地区字段，并给出可读包邮结论。"""

from __future__ import annotations

import re
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"D:\Users\shenghua\Downloads\商品筛选\douhuo_scraper")
from client import DouhuoClient  # noqa: E402
from excelio import FIELD_ZH, flatten_value, union_keys  # noqa: E402

PROFILE = Path(r"D:\Users\shenghua\Downloads\商品筛选\state\browser_profile")
SRC = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\油类\油类全量商品_京东类目版.xlsx")
OUT = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\油类\油类全量商品_京东类目版_含包邮地区.xlsx")
STATUS = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\progress_status.txt")
CACHE = Path(r"D:\Users\shenghua\Downloads\商品筛选\output\油类\freight_cache.json")

HEADER_FILL = PatternFill("solid", fgColor="1D1F27")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LABEL_FILL = PatternFill("solid", fgColor="3D4450")
LABEL_FONT = Font(color="FFFFFF", size=10)

PREFERRED = [
    "子类", "子类细分", "品牌", "规格",
    "goods_id", "spu_sn", "spu_name", "default_sku",
    "supply_type", "supply_name", "plat_price", "retail_price", "profit",
    "包邮结论", "is_freight", "不包邮_不发货地区", "包邮地区",
    "min_reduce", "max_reduce", "status",
    "main_img", "商品链接", "匹配关键词", "采集时间",
]
ZH_EXTRA = {
    "包邮结论": "可读结论 freight_summary",
    "is_freight": "原始值 1=包邮 2=不包邮/按地区",
    "不包邮_不发货地区": "接口 no_deliver",
    "包邮地区": "接口 free_shipping（仅特定包邮时有值）",
    "子类": "标准子类(京东/淘宝)",
    "子类细分": "细分(浓香/低芥酸等)",
}

LEGEND = [
    ("包邮结论", "规则："),
    ("", "1) is_freight=1 且 无地区限制 → 全国包邮（不包邮地区无）"),
    ("", "2) is_freight=1 但 no_deliver 有值 → 除所列地区外包邮"),
    ("", "3) is_freight=2 且 free_shipping 有值 → 仅所列地区包邮，其余不包邮"),
    ("", "4) is_freight=2 且 no_deliver 有值 → 以下地区不包邮/不发货"),
    ("", "5) is_freight=2 且两者皆空 → 不包邮（接口未返回地区明细，可能按地址计运费）"),
    ("is_freight", "1=包邮；2=不包邮或按地区"),
    ("不包邮_不发货地区", "detail.no_deliver，逗号分隔省市"),
    ("包邮地区", "free_shipping 接口，仅当明确写了包邮地区时有值"),
]


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def write_status(msg: str) -> None:
    try:
        STATUS.parent.mkdir(parents=True, exist_ok=True)
        STATUS.write_text(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n", encoding="utf-8")
    except OSError:
        pass


def norm_regions(raw) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s or s.lower() in {"null", "none", "无"}:
        return ""
    s = s.strip("-")
    parts = [p for p in re.split(r"[-,，、;；/|]+", s) if p and p != "无"]
    return "、".join(parts)


def summarize(is_freight, no_deliver: str, free_ship: str) -> str:
    nd = (no_deliver or "").strip()
    fs = (free_ship or "").strip()
    freight = is_freight
    try:
        freight = int(is_freight) if is_freight not in (None, "") else None
    except (TypeError, ValueError):
        freight = None

    if freight == 1 and not nd:
        return "全国包邮（不包邮地区无）"
    if freight == 1 and nd:
        return f"包邮，但以下地区除外：{nd}"
    if freight == 2 and fs and not nd:
        return f"仅以下地区包邮：{fs}；其余地区不包邮"
    if freight == 2 and nd and not fs:
        return f"不包邮/限运地区：{nd}"
    if freight == 2 and nd and fs:
        return f"包邮地区：{fs}；不包邮/不发货：{nd}"
    if freight == 2:
        return "不包邮（接口未返回地区明细，可能按收货地址计运费）"
    if nd:
        return f"限运地区：{nd}"
    if fs:
        return f"包邮地区：{fs}"
    return "未知（无运费地区信息）"


def dump_sheet(ws, rows, headers):
    for col, key in enumerate(headers, start=1):
        c1 = ws.cell(1, col, key)
        c1.fill = HEADER_FILL
        c1.font = HEADER_FONT
        c2 = ws.cell(2, col, ZH_EXTRA.get(key) or FIELD_ZH.get(key) or key)
        c2.fill = LABEL_FILL
        c2.font = LABEL_FONT
        ws.column_dimensions[get_column_letter(col)].width = min(max(len(key) + 2, 12), 30)
    for r_i, row in enumerate(rows, start=3):
        for c_i, key in enumerate(headers, start=1):
            ws.cell(r_i, c_i, flatten_value(row.get(key)))
    ws.freeze_panes = "A3"


def main() -> int:
    import json

    wb = load_workbook(SRC, data_only=True)
    ws = wb["食用油全量"]
    raw = list(ws.iter_rows(values_only=True))
    headers_in = list(raw[0])
    idx = {h: i for i, h in enumerate(headers_in)}
    base_rows = []
    for r in raw[2:]:
        if not r or r[idx["goods_id"]] in (None, ""):
            continue
        row = {h: (r[idx[h]] if h in idx and idx[h] < len(r) else None) for h in headers_in}
        base_rows.append(row)
    wb.close()
    log(f"读入 {len(base_rows)} 行，开始补运费地区")

    cache: dict[str, dict] = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
            log(f"缓存命中 {len(cache)}")
        except json.JSONDecodeError:
            cache = {}

    t0 = time.time()
    last = t0
    with DouhuoClient(PROFILE, headless=True) as client:
        log("登录态 " + client.login_state())
        client._page.wait_for_timeout(3000)
        for i, row in enumerate(base_rows, 1):
            gid = str(row.get("goods_id") or "")
            if gid in cache:
                info = cache[gid]
            else:
                no_deliver = ""
                free_ship = ""
                detail_freight = row.get("is_freight")
                try:
                    d = client.post("/manager/Choice/getGoodsDetail", {"goods_id": gid}) or {}
                    if isinstance(d, dict):
                        no_deliver = norm_regions(d.get("no_deliver"))
                        if d.get("is_freight") not in (None, ""):
                            detail_freight = d.get("is_freight")
                except Exception as exc:  # noqa: BLE001
                    log(f"  detail fail {gid}: {exc}")
                try:
                    fs = client.post("/manager/Choice/free_shipping", {"goods_id": gid})
                    free_ship = norm_regions(fs)
                except Exception as exc:  # noqa: BLE001
                    log(f"  free_shipping fail {gid}: {exc}")
                info = {
                    "is_freight": detail_freight,
                    "no_deliver": no_deliver,
                    "free_shipping": free_ship,
                }
                cache[gid] = info
                if i % 50 == 0:
                    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

            row["is_freight"] = info.get("is_freight")
            row["不包邮_不发货地区"] = info.get("no_deliver") or ""
            row["包邮地区"] = info.get("free_shipping") or ""
            row["包邮结论"] = summarize(
                info.get("is_freight"), info.get("no_deliver") or "", info.get("free_shipping") or ""
            )

            now = time.time()
            if now - last >= 300 or i == len(base_rows) or i % 200 == 0:
                elapsed = now - t0
                rate = i / elapsed if elapsed else 0
                eta = (len(base_rows) - i) / rate / 60 if rate else 0
                msg = f"[进度] 运费补全 {i}/{len(base_rows)} ({i*100//len(base_rows)}%) 已用{elapsed/60:.1f}分钟 ETA≈{eta:.0f}分钟"
                log(msg)
                write_status(msg)
                last = now
            elif i % 40 == 0:
                log(f"  … {i}/{len(base_rows)}")

    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    headers = union_keys(base_rows, PREFERRED)
    out_wb = Workbook()
    ws0 = out_wb.active
    ws0.title = "字段说明"
    ws0.append(["字段 Field", "说明 Description"])
    for c in ws0[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for a, b in LEGEND:
        ws0.append([a, b])
    for h in headers:
        ws0.append([h, ZH_EXTRA.get(h) or FIELD_ZH.get(h) or h])
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 70

    dump_sheet(out_wb.create_sheet("食用油全量"), base_rows, headers)
    by = {}
    for r in base_rows:
        by.setdefault(r.get("子类") or "未分类", []).append(r)
    for sub, items in sorted(by.items(), key=lambda x: -len(x[1])):
        dump_sheet(out_wb.create_sheet(re.sub(r'[\\/*?:\[\]]', "_", sub)[:31]), items, headers)

    # 包邮结论分布
    from collections import Counter
    dist = Counter(r.get("包邮结论") for r in base_rows)
    ws_s = out_wb.create_sheet("包邮结论分布")
    ws_s.append(["包邮结论", "商品数"])
    for c in ws_s[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for k, v in dist.most_common():
        ws_s.append([k, v])
    ws_s.column_dimensions["A"].width = 80

    out_wb.save(OUT)
    log("=" * 40)
    log(f"完成 {len(base_rows)} 条 → {OUT}")
    for k, v in dist.most_common(12):
        log(f"  {v:4d}  {k}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
