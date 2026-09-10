"""JSONL → Excel 导出：双行表头（英文字段 + 中文说明）+ 全字段保留。"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1D1F27")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LABEL_FILL = PatternFill("solid", fgColor="3D4450")
LABEL_FONT = Font(color="FFFFFF", bold=False, size=10)

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")

# 接口字段 → 中文表头（第二行）。未列出的字段用原名。
FIELD_ZH: dict[str, str] = {
    "_row_type": "行类型",
    "goods_id": "商品ID",
    "id": "商品ID(详情)",
    "spu_sn": "商品编码",
    "goods_sn": "商品货号",
    "goods_sign": "商品标识",
    "spu_name": "商品名称",
    "goods_name": "商品名称(轮播)",
    "supply_type": "渠道类型",
    "supply_name": "渠道名称",
    "plat_price": "代发价",
    "retail_price": "零售价",
    "market_price": "市场价",
    "collection_price": "集采价",
    "profit": "利润率(%)",
    "min_reduce": "集采价下限",
    "max_reduce": "集采价上限",
    "is_freight": "是否包邮",
    "status": "上下架状态",
    "sort": "推荐指数",
    "sort_init": "初始排序",
    "sort_vari": "销量排序",
    "sort_set": "人工排序",
    "sort_total": "综合排序",
    "default_sku": "默认SKU",
    "main_img": "主图",
    "cover_img": "封面图",
    "top_img": "顶部图",
    "bottom_img": "底部图",
    "detail_img": "详情图",
    "is_select": "是否已选",
    "is_test": "是否测试",
    "is_update": "是否更新中",
    "update_time": "更新时间",
    "create_time": "创建时间",
    "sync_time": "同步时间",
    "audit_time": "审核时间",
    "compute_time": "计算时间",
    "cate_pid": "父类目ID",
    "cate_id": "类目ID",
    "cate_name": "类目名称",
    "brand_id": "品牌ID",
    "brand_name": "品牌名称",
    "tag_id": "标签ID",
    "tax_id": "税率ID",
    "goods_type": "商品类型",
    "supplier_id": "供应商ID",
    "supplier_name": "供应商名称",
    "sub_supplier_id": "子供应商ID",
    "sub_supplier_name": "子供应商名称",
    "limit_num": "限购数",
    "lowest_num": "起批量",
    "is_return": "是否可退",
    "is_display": "是否展示",
    "no_deliver": "不发货地区",
    "region": "地区",
    "sku_list": "SKU列表",
    "explain": "说明",
    "expense": "费用",
    "table_name": "表名",
    "sn": "编号",
    "old_sn": "原编号",
    "c_ext_link": "电脑外链",
    "m_ext_link": "手机外链",
    "stats_sale_num": "销量",
    "min_price_value": "最低代发价",
    "min_retail_value": "最低零售价",
    "min_profit_value": "最低利润率",
    "max_price_value": "最高代发价",
    "max_retail_value": "最高零售价",
    "max_profit_value": "最高利润率",
    "商品链接": "商品链接",
    "nav_id": "专区ID",
    "nav_name": "专区名称",
    "nav_type": "专区类型",
    "nav_icon": "专区图标",
    "icon_type": "图标类型",
    "version": "版本",
    "link_url": "跳转链接",
    "mobile_url": "手机跳转链接",
    "bg_color": "背景色",
    "base_map": "底图",
    "is_auth": "是否授权",
    "goods_count": "内含商品数",
    "img_id": "图片ID",
    "img_name": "图片名称",
    "img_path": "图片地址",
    "mobile_img_path": "手机图片地址",
    "type": "图片类型",
    "解析nav_id": "链接解析专区ID",
    "解析goods_id": "链接解析商品ID",
    "解析nav_type": "链接解析专区类型",
    "采集时间": "采集时间",
    "来源分区": "来源分区",
    "来源banner_id": "来源轮播图ID",
    "来源banner名称": "来源轮播图名称",
    "来源banner链接": "来源轮播图链接",
    "total": "总条数",
    "pages": "总页数",
    "page": "页码",
    "limit": "每页条数",
}


def zh_label(key: str) -> str:
    return FIELD_ZH.get(key, key)


def flatten_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return _ILLEGAL_XML.sub("", value)
    return value


def union_keys(rows: Iterable[dict], preferred: list[str] | None = None) -> list[str]:
    seen: list[str] = list(preferred or [])
    known = set(seen)
    for row in rows:
        for key in row.keys():
            if key not in known:
                known.add(key)
                seen.append(key)
    return seen


def write_excel(path: Path, rows: list[dict], preferred: list[str] | None = None) -> int:
    """双行表头：第1行英文字段键，第2行中文说明；数据从第3行开始。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = union_keys(rows, preferred)
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    for col, name in enumerate(headers, start=1):
        zh = zh_label(name)
        c1 = ws.cell(row=1, column=col, value=name)
        c1.fill = HEADER_FILL
        c1.font = HEADER_FONT
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2 = ws.cell(row=2, column=col, value=zh)
        c2.fill = LABEL_FILL
        c2.font = LABEL_FONT
        c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = max(len(name) + 2, len(zh) * 2 + 2, 12)
        ws.column_dimensions[get_column_letter(col)].width = min(width, 36)
    for r_idx, row in enumerate(rows, start=3):
        for c_idx, key in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c_idx, value=flatten_value(row.get(key)))
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 22
    wb.save(path)
    return len(rows)


GOODS_PREFERRED = [
    "goods_id", "spu_sn", "spu_name", "supply_type", "supply_name",
    "plat_price", "retail_price", "profit", "min_reduce", "max_reduce",
    "is_freight", "status", "sort", "default_sku", "main_img",
    "is_select", "is_test", "is_update", "update_time",
    "商品链接", "nav_id", "nav_name", "采集时间", "来源分区",
]


def normalize_goods_row(raw: dict, *, nav_id: Any = "", nav_name: str = "", source: str = "") -> dict:
    goods_id = raw.get("goods_id") or raw.get("id") or ""
    row = dict(raw)
    row["nav_id"] = nav_id
    row["nav_name"] = nav_name
    row["来源分区"] = source
    row["商品链接"] = (
        f"https://www.douhuomall.com/choice/v3/#/goods?id={goods_id}&nav_type=1"
        if goods_id != "" else ""
    )
    row["采集时间"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return row


BANNER_PREFERRED = [
    "_row_type",
    "goods_id", "spu_sn", "spu_name", "goods_name", "supply_type", "supply_name",
    "plat_price", "retail_price", "profit", "min_reduce", "max_reduce",
    "is_freight", "status", "sort", "default_sku", "main_img", "cover_img",
    "商品链接", "nav_id", "nav_name",
    "来源banner_id", "来源banner名称", "来源banner链接",
    "img_id", "img_name", "img_path", "mobile_img_path", "link_url", "mobile_url", "type",
    "解析nav_id", "解析goods_id", "解析nav_type", "采集时间", "来源分区",
]


def normalize_banner_row(raw: dict, parsed: dict) -> dict:
    row = dict(raw)
    row["解析nav_id"] = parsed.get("id") or parsed.get("nav_id") or ""
    row["解析goods_id"] = parsed.get("goods_id") or ""
    row["解析nav_type"] = parsed.get("nav_type") or ""
    row["采集时间"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return row
